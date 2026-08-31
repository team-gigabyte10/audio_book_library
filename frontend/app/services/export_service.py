import os
import json
import csv
from typing import Dict, Any, List
from docx import Document
from docx.shared import Pt


class ExportService:
    """
    Generates downloadable export files (.txt, .json, .docx, .csv) for processed books.
    """

    def __init__(self, export_dir: str):
        self.export_dir = export_dir
        os.makedirs(self.export_dir, exist_ok=True)

    def export_txt(self, job_data: Dict[str, Any]) -> str:
        """
        Exports extracted text to a .txt file.
        """
        filename = f"{job_data['job_id']}_extracted.txt"
        file_path = os.path.join(self.export_dir, filename)

        pages = job_data.get("pages", [])
        content_lines = [
            f"# {job_data.get('original_filename', 'Bangla Book')}",
            f"# Total Pages: {job_data.get('total_pages', 0)}",
            "=" * 50,
            ""
        ]

        for p in sorted(pages, key=lambda x: x["page_number"]):
            content_lines.append(f"--- PAGE {p['page_number']} ---")
            text = p.get("cleaned_text") or p.get("raw_text") or ""
            content_lines.append(text)
            content_lines.append("")

        with open(file_path, "w", encoding="utf-8") as f:
            f.write("\n".join(content_lines))

        return file_path

    def export_json(self, job_data: Dict[str, Any]) -> str:
        """
        Exports full job metadata and page-level results to a .json file.
        """
        filename = f"{job_data['job_id']}_metadata.json"
        file_path = os.path.join(self.export_dir, filename)

        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(job_data, f, ensure_ascii=False, indent=2)

        return file_path

    def export_docx(self, job_data: Dict[str, Any]) -> str:
        """
        Exports document to Microsoft Word .docx format with chapter structure and paragraphs.
        """
        filename = f"{job_data['job_id']}_extracted.docx"
        file_path = os.path.join(self.export_dir, filename)

        doc = Document()

        # Add Title
        title_heading = doc.add_heading(job_data.get("original_filename", "Bangla Book Extracted Text"), level=0)
        
        chapters = job_data.get("chapters", [])
        pages = sorted(job_data.get("pages", []), key=lambda x: x["page_number"])
        page_dict = {p["page_number"]: p for p in pages}

        if chapters:
            for ch in chapters:
                doc.add_heading(ch["title"], level=1)
                for p_num in range(ch["start_page"], ch["end_page"] + 1):
                    if p_num in page_dict:
                        text = page_dict[p_num].get("cleaned_text") or page_dict[p_num].get("raw_text") or ""
                        if text:
                            for paragraph in text.split("\n\n"):
                                p = doc.add_paragraph(paragraph)
                                p.style.font.name = "Kalpurush"  # Common Bengali Word font
        else:
            for p in pages:
                text = p.get("cleaned_text") or p.get("raw_text") or ""
                if text:
                    for paragraph in text.split("\n\n"):
                        p_elem = doc.add_paragraph(paragraph)

        doc.save(file_path)
        return file_path

    def export_csv(self, job_data: Dict[str, Any]) -> str:
        """
        Exports page-level results to CSV format.
        """
        filename = f"{job_data['job_id']}_extracted.csv"
        file_path = os.path.join(self.export_dir, filename)

        pages = sorted(job_data.get("pages", []), key=lambda x: x["page_number"])

        fieldnames = ["page_number", "processing_method", "page_type", "character_count", "word_count", "ocr_used", "text"]

        with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for p in pages:
                writer.writerow({
                    "page_number": p.get("page_number"),
                    "processing_method": p.get("processing_method"),
                    "page_type": p.get("page_type"),
                    "character_count": p.get("character_count", 0),
                    "word_count": p.get("word_count", 0),
                    "ocr_used": p.get("ocr_used", False),
                    "text": p.get("cleaned_text") or p.get("raw_text") or ""
                })

        return file_path

    def export_md(self, job_data: Dict[str, Any]) -> str:
        """
        Exports extracted text to a structured, clean Markdown (.md) file with headers, TOC, lists, and blockquotes.
        """
        import re
        from app.services.font_converter import font_converter

        filename = f"{job_data['job_id']}_extracted.md"
        file_path = os.path.join(self.export_dir, filename)

        orig_filename = job_data.get("original_filename", "Bangla Book")
        title = orig_filename.rsplit(".", 1)[0].replace("_", " ")

        pages = sorted(job_data.get("pages", []), key=lambda x: x["page_number"])
        chapters = job_data.get("chapters", [])

        md_lines = []

        # Document Header
        md_lines.append(f"# {title}")
        md_lines.append(f"*(বাংলা বই টেক্সট সারসংক্ষেপ / Extracted Summary)*\n")
        md_lines.append(f"**মোট পৃষ্ঠা (Total Pages):** {job_data.get('total_pages', len(pages))}  ")
        md_lines.append(f"**প্রসেসিং স্ট্যাটাস:** {job_data.get('status', 'COMPLETED')}  \n")
        md_lines.append("---")
        md_lines.append("")

        # Table of Contents Header
        if chapters:
            md_lines.append("## সূচিপত্র (Table of Contents)")
            md_lines.append("")
            for idx, ch in enumerate(chapters, 1):
                clean_title = font_converter.convert(ch["title"])
                anchor = re.sub(r"[^\w\s-]", "", clean_title.lower()).replace(" ", "-")
                md_lines.append(f"{idx}. [{clean_title}](#{anchor})")
            md_lines.append("")
        # Process page content
        for p in pages:
            text = p.get("cleaned_text") or p.get("raw_text") or ""
            if not text.strip():
                continue

            page_num = p["page_number"]
            cleaned_text = font_converter.convert(text)

            md_lines.append(f"<!-- Page {page_num} -->")

            # Check if this page is purely a Table of Contents page
            is_toc_page = "সূচিপত্র" in cleaned_text and page_num <= 3

            raw_paragraphs = [para.strip() for para in cleaned_text.split("\n\n") if para.strip()]

            for para in raw_paragraphs:
                lines = [l.strip() for l in para.split("\n") if l.strip()]
                if not lines:
                    continue

                first_line = lines[0]

                # 1. Numbered TOC list items on TOC page -> Keep as numbered list
                if is_toc_page and re.match(r"^[০-৯0-9]+[\.\)]", first_line):
                    for line in lines:
                        md_lines.append(line)
                    md_lines.append("")
                    continue

                # 2. Main Chapter Heading detection (generic chapter patterns)
                is_main_heading = (
                    re.match(r"^(?:পাঠ|অধ্যায়|অধ্যায়|পর্ব|ভূমিকা|উপসংহার|সূচিপত্র|Chapter)\s*[-—:\s]*[০-৯0-9]*", first_line, re.IGNORECASE) or
                    any(first_line.startswith(h) for h in ["ভূমিকা:", "উপসংহার:", "সূচিপত্র"])
                )

                # 3. Subheading detection (generic structural rules - zero hardcoded book words)
                is_subheading = (
                    (first_line.endswith(":") and len(first_line) < 60 and not first_line.startswith(("বলতেন:", "বলতেনঃ", "বলেন:"))) or
                    (re.match(r"^[০-৯0-9]+\.\s+[^\।]{3,40}$", first_line) and not is_toc_page) or
                    (len(first_line) < 40 and not first_line.endswith("।") and not first_line.endswith("?") and not first_line.endswith("!"))
                )

                if is_main_heading:
                    md_lines.append(f"## {first_line}\n")
                    # Append remaining lines of paragraph
                    for extra in lines[1:]:
                        md_lines.append(extra)
                    md_lines.append("")
                elif is_subheading:
                    md_lines.append(f"### {first_line}\n")
                    for extra in lines[1:]:
                        md_lines.append(extra)
                    md_lines.append("")
                elif first_line.startswith(("•", "-", "*")):
                    # Bullet point paragraph: merge soft breaks within bullet item
                    bullet_text = " ".join([l.lstrip("•-* ").strip() for l in lines])
                    md_lines.append(f"- {bullet_text}\n")
                elif "বলতেন:" in first_line or '"' in first_line or '“' in first_line or '”' in first_line:
                    # Quote block
                    quote_text = " ".join(lines)
                    md_lines.append(f"> {quote_text}\n")
                else:
                    # Regular text paragraph
                    para_text = " ".join(lines)
                    md_lines.append(f"{para_text}\n")

            md_lines.append("---\n")

        with open(file_path, "w", encoding="utf-8") as f:
            f.write("\n".join(md_lines))

        return file_path


